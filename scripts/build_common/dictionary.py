from __future__ import annotations

import csv
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = [
    "Variável",
    "Nome Descritivo",
    "Tipo de Dado",
    "Unidade de Medida",
    "Domínio / Categoria",
    "Descrição / Significado",
    "Observações",
]

SHEETS = (
    ("Base bruta", "DicionarioBaseBruta"),
    ("Variáveis derivadas", "DicionarioVariaveisDerivadas"),
)


class DictionarySourceError(ValueError):
    """Indica que uma fonte CSV do dicionário é inválida."""


def read_dictionary_source(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise DictionarySourceError(f"Fonte do dicionário não encontrada: {path}")

    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != HEADERS:
            raise DictionarySourceError(
                f"Cabeçalho inválido em {path}. Esperado: {HEADERS}. Recebido: {reader.fieldnames}."
            )
        rows = []
        seen: set[str] = set()
        for line_number, row in enumerate(reader, start=2):
            normalized = {header: (row.get(header) or "").strip() for header in HEADERS}
            if not any(normalized.values()):
                continue
            variable = normalized["Variável"]
            if not variable:
                raise DictionarySourceError(f"Variável vazia em {path}:{line_number}")
            if variable in seen:
                raise DictionarySourceError(f"Variável duplicada em {path}:{line_number}: {variable}")
            if not normalized["Descrição / Significado"]:
                raise DictionarySourceError(
                    f"Descrição ausente em {path}:{line_number} para a variável {variable}"
                )
            seen.add(variable)
            rows.append(normalized)

    if not rows:
        raise DictionarySourceError(f"Nenhuma variável documentada em {path}")
    return rows


def _style_sheet(sheet, rows: list[dict[str, str]], table_name: str) -> None:
    dark_blue = "1F4E78"
    pale_blue = "EAF2F8"
    light_gray = "D9D9D9"
    white = "FFFFFF"
    body_font = Font(name="Arial", size=10, color="000000")
    header_font = Font(name="Arial", size=10, bold=True, color=white)
    thin = Side(style="thin", color=light_gray)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{len(rows) + 1}"
    sheet.row_dimensions[1].height = 30
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"

    widths = [24, 38, 22, 20, 44, 64, 72]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_index in range(2, len(rows) + 2):
        estimated_lines = 1
        for column_index in range(1, 8):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            value = str(cell.value or "")
            # Aproxima a quebra de linha do Excel de acordo com a largura de
            # cada coluna, incluindo quebras explícitas presentes na fonte.
            usable_chars = max(8, int(widths[column_index - 1] * 1.25))
            line_count = sum(max(1, math.ceil(len(part) / usable_chars)) for part in value.splitlines())
            estimated_lines = max(estimated_lines, line_count)
        sheet.row_dimensions[row_index].height = min(150, max(28, 15 * estimated_lines))

    table = Table(displayName=table_name, ref=f"A1:G{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def build_dictionary(raw_source: Path, derived_source: Path, output: Path) -> Path:
    source_rows = (
        read_dictionary_source(raw_source),
        read_dictionary_source(derived_source),
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Dicionário de dados do projeto de atrasos de voos"
    workbook.properties.subject = "Base bruta e variáveis derivadas"
    workbook.properties.creator = "Grupo da disciplina de Introdução ao Machine Learning"

    for (sheet_name, table_name), rows in zip(SHEETS, source_rows, strict=True):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(HEADERS)
        for row in rows:
            sheet.append([row[header] for header in HEADERS])
        _style_sheet(sheet, rows, table_name)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def validate_dictionary(output: Path, expected_counts: tuple[int, int]) -> None:
    if not output.is_file() or output.stat().st_size == 0:
        raise DictionarySourceError(f"Dicionário não foi gerado: {output}")
    workbook = load_workbook(output, read_only=False, data_only=False)
    expected_sheets = [name for name, _ in SHEETS]
    if workbook.sheetnames != expected_sheets:
        raise DictionarySourceError(
            f"Abas inválidas. Esperado: {expected_sheets}. Recebido: {workbook.sheetnames}."
        )
    for sheet_name, expected_count in zip(expected_sheets, expected_counts, strict=True):
        sheet = workbook[sheet_name]
        headers = [sheet.cell(1, column).value for column in range(1, 8)]
        if headers != HEADERS:
            raise DictionarySourceError(f"Cabeçalho inválido na aba {sheet_name}: {headers}")
        if sheet.max_row != expected_count + 1:
            raise DictionarySourceError(
                f"Quantidade de linhas inválida em {sheet_name}: {sheet.max_row - 1}"
            )
        if sheet.freeze_panes != "A2":
            raise DictionarySourceError(f"Cabeçalho não congelado na aba {sheet_name}")
        if sheet.auto_filter.ref != f"A1:G{expected_count + 1}":
            raise DictionarySourceError(f"Autofiltro inválido na aba {sheet_name}")
        if len(sheet.tables) != 1:
            raise DictionarySourceError(f"Tabela estruturada ausente na aba {sheet_name}")

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INCLUDE_PATTERN = re.compile(r"^\{\{\s*include\s+(.+?)\s*\}\}$")
INLINE_PATTERN = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)")
PLACEHOLDER_PATTERN = re.compile(r"\{\{.+?\}\}")


class ReportSourceError(ValueError):
    """Indica que a fonte Markdown do relatório é inválida."""


def expand_markdown(path: Path, allowed_root: Path, stack: tuple[Path, ...] = ()) -> str:
    resolved = path.resolve()
    allowed = allowed_root.resolve()
    if allowed not in resolved.parents and resolved != allowed:
        raise ReportSourceError(f"Inclusão fora de sources/: {path}")
    if resolved in stack:
        chain = " -> ".join(str(item) for item in (*stack, resolved))
        raise ReportSourceError(f"Ciclo de inclusão no Markdown: {chain}")
    if not resolved.is_file():
        raise ReportSourceError(f"Fonte Markdown não encontrada: {resolved}")

    output: list[str] = []
    for line in resolved.read_text(encoding="utf-8").splitlines():
        match = INCLUDE_PATTERN.match(line.strip())
        if match:
            child = (resolved.parent / match.group(1)).resolve()
            output.append(expand_markdown(child, allowed, (*stack, resolved)).rstrip())
        else:
            output.append(line)
    return "\n".join(output).strip() + "\n"


def _set_font(run, *, name: str = "Times New Roman", size: int = 12, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_style(document: Document, name: str, *, size: int, bold: bool, alignment=None):
    styles = document.styles
    if name not in styles:
        style_type = WD_STYLE_TYPE.PARAGRAPH
        styles.add_style(name, style_type)
    style = styles[name]
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    if alignment is not None:
        style.paragraph_format.alignment = alignment
    return style


def configure_styles(document: Document) -> None:
    normal = _configure_style(
        document, "Normal", size=12, bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    title = _configure_style(
        document, "Title", size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER
    )
    title.paragraph_format.space_after = Pt(14)

    heading1 = _configure_style(
        document, "Heading 1", size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT
    )
    heading1.paragraph_format.space_before = Pt(12)
    heading1.paragraph_format.space_after = Pt(6)
    heading1.paragraph_format.keep_with_next = True

    heading2 = _configure_style(
        document, "Heading 2", size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT
    )
    heading2.paragraph_format.space_before = Pt(10)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True

    list_style = _configure_style(
        document, "List Bullet", size=12, bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    list_style.paragraph_format.left_indent = Inches(0.35)
    list_style.paragraph_format.first_line_indent = Inches(-0.2)
    list_style.paragraph_format.line_spacing = 1.3
    list_style.paragraph_format.space_after = Pt(4)


def _add_inline(paragraph, text: str) -> None:
    for token in filter(None, INLINE_PATTERN.split(text)):
        # None preserva a formatação herdada do estilo do parágrafo, o que é
        # essencial para títulos e cabeçalhos continuarem em negrito.
        bold = None
        italic = None
        value = token
        if token.startswith("***") and token.endswith("***"):
            value = token[3:-3]
            bold = True
            italic = True
        elif token.startswith("**") and token.endswith("**"):
            value = token[2:-2]
            bold = True
        elif token.startswith("*") and token.endswith("*"):
            value = token[1:-1]
            italic = True
        run = paragraph.add_run(value)
        _set_font(run, bold=bold, italic=italic)


def build_report(source: Path, template: Path, output: Path, sources_root: Path) -> Path:
    if not template.is_file():
        raise ReportSourceError(f"Template institucional não encontrado: {template}")
    markdown = expand_markdown(source, sources_root)
    unresolved = [line for line in markdown.splitlines() if PLACEHOLDER_PATTERN.search(line)]
    if unresolved:
        raise ReportSourceError(f"Diretivas ou placeholders não resolvidos: {unresolved[:3]}")

    document = Document(template)
    configure_styles(document)
    document.core_properties.title = "Projeto parcial da Avaliação 1"
    document.core_properties.subject = "Introdução, objetivos, referencial teórico e referências"
    document.core_properties.creator = "Grupo da disciplina de Introdução ao Machine Learning"

    current_section = ""
    cover_mode = True
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            _add_inline(paragraph, line[4:])
            current_section = line[4:]
            cover_mode = False
        elif line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            _add_inline(paragraph, line[3:])
            current_section = line[3:]
            cover_mode = False
        elif line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            _add_inline(paragraph, line[2:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            bullet = paragraph.add_run("• ")
            _set_font(bullet)
            _add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph(style="Normal")
            _add_inline(paragraph, line)
            if cover_mode:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(6)
            elif current_section == "4. Referências":
                # URLs e identificadores longos produzem espaçamentos excessivos
                # quando referências bibliográficas são justificadas.
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.left_indent = Inches(0.4)
                paragraph.paragraph_format.first_line_indent = Inches(-0.4)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(6)

    if not any(p.text.strip().startswith("1. Introdução") for p in document.paragraphs):
        raise ReportSourceError("Seção '1. Introdução' ausente do relatório")
    if not any(p.text.strip().startswith("4. Referências") for p in document.paragraphs):
        raise ReportSourceError("Seção '4. Referências' ausente do relatório")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def validate_report(output: Path) -> None:
    if not output.is_file() or output.stat().st_size == 0:
        raise ReportSourceError(f"Relatório não foi gerado: {output}")
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required = ("1. Introdução", "2. Objetivos", "3. Referencial Teórico", "4. Referências")
    missing = [section for section in required if section not in text]
    if missing:
        raise ReportSourceError(f"Seções ausentes no relatório: {missing}")
    if PLACEHOLDER_PATTERN.search(text):
        raise ReportSourceError("O relatório contém placeholder não substituído")
    if not document.sections:
        raise ReportSourceError("O relatório não possui configuração de seção")
    section = document.sections[0]
    if not any(paragraph.text.strip() or paragraph._p.xpath(".//w:drawing") for paragraph in section.header.paragraphs):
        raise ReportSourceError("Cabeçalho institucional ausente")
    if not any(paragraph.text.strip() or paragraph._p.xpath(".//w:drawing") for paragraph in section.footer.paragraphs):
        raise ReportSourceError("Rodapé institucional ausente")

from __future__ import annotations

import csv
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_common.dictionary import (  # noqa: E402
    ALLOWED_DATA_TYPES,
    DictionarySourceError,
    SHEET_NAME,
    build_dictionary,
    read_dictionary_source,
)


class DictionarySourcesTest(unittest.TestCase):
    def test_expected_source_counts(self) -> None:
        base = read_dictionary_source(ROOT / "sources/dicionario/base_consolidada.csv")
        analytical = read_dictionary_source(ROOT / "sources/dicionario/variaveis_analiticas.csv")
        self.assertEqual(25, len(base))
        self.assertEqual(20, len(analytical))
        self.assertIn("codigo_faixa_atraso", {row["Variável"] for row in analytical})

    def test_sources_use_only_standardized_data_types(self) -> None:
        for source_name in ("base_consolidada.csv", "variaveis_analiticas.csv"):
            rows = read_dictionary_source(ROOT / "sources/dicionario" / source_name)
            observed_types = {row["Tipo de Dado"] for row in rows}
            self.assertTrue(observed_types)
            self.assertLessEqual(observed_types, ALLOWED_DATA_TYPES)

    def test_base_dictionary_matches_the_csv_header(self) -> None:
        dictionary = read_dictionary_source(ROOT / "sources/dicionario/base_consolidada.csv")
        dictionary_names = [row["Variável"] for row in dictionary]
        archive_path = ROOT / "data/raw/base_lucimar_nascimento_v2.zip"
        with zipfile.ZipFile(archive_path) as archive:
            csv_names = [name for name in archive.namelist() if name.lower().endswith(".csv")]
            self.assertEqual(["base_lucimar_nascimento_v2.csv"], csv_names)
            with archive.open(csv_names[0]) as raw:
                text = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
                csv_header = next(csv.reader(text))
        self.assertEqual(csv_header, dictionary_names)

    def test_generated_dictionary_contains_only_the_delivered_base(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "dictionary.xlsx"
            build_dictionary(ROOT / "sources/dicionario/base_consolidada.csv", output)
            workbook = load_workbook(output, read_only=False, data_only=False)
            self.assertEqual([SHEET_NAME], workbook.sheetnames)
            sheet = workbook[SHEET_NAME]
            self.assertEqual(26, sheet.max_row)
            self.assertEqual(7, sheet.max_column)
            self.assertEqual("companhia_icao", sheet["A2"].value)
            self.assertEqual("linha_origem", sheet["A26"].value)
            workbook.close()

    def test_duplicate_variable_is_rejected(self) -> None:
        temporary = ROOT / "target/test_duplicate_dictionary.csv"
        temporary.parent.mkdir(parents=True, exist_ok=True)
        source = (ROOT / "sources/dicionario/base_consolidada.csv").read_text(encoding="utf-8")
        lines = source.splitlines()
        temporary.write_text("\n".join([*lines, lines[1]]) + "\n", encoding="utf-8")
        try:
            with self.assertRaises(DictionarySourceError):
                read_dictionary_source(temporary)
        finally:
            temporary.unlink(missing_ok=True)

    def test_invalid_header_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "invalid.csv"
            source.write_text("Variavel,Descricao\ncoluna,Texto\n", encoding="utf-8")
            with self.assertRaisesRegex(DictionarySourceError, "Cabeçalho inválido"):
                read_dictionary_source(source)

    def test_nonstandard_data_type_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "invalid_type.csv"
            source.write_text(
                "Variável,Nome Descritivo,Tipo de Dado,Unidade de Medida,"
                "Domínio / Categoria,Descrição / Significado,Observações\n"
                "coluna,Coluna,Categórico (texto),Não aplicável,Texto,Descrição,\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(DictionarySourceError, "Tipo de dado inválido"):
                read_dictionary_source(source)


if __name__ == "__main__":
    unittest.main()
